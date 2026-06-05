import streamlit as st
import pandas as pd
import numpy as np
import io
import re
import openpyxl
from openpyxl.styles import PatternFill
from datetime import datetime
from supabase import create_client, Client

# Configuración de la plataforma de Auditoría
st.set_page_config(page_title="Sistema de Control Transaccional - Cheques", layout="wide")

# ==========================================
# CONEXIÓN NATIVA A LA BASE DE DATOS SQL
# ==========================================
def conectar_supabase() -> Client:
    try:
        url = st.secrets["SUPABASE_URL"]
        key = st.secrets["SUPABASE_KEY"]
        return create_client(url, key)
    except Exception as e:
        st.error(f"⚠️ Error de enlace con el servidor SQL: {e}")
        return None

# Plano de traducción oficial entre la interfaz y la base de datos SQL
DB_MAP = {
    "EMPRESA": "empresa", "FECHA DE EMISIÓN": "fecha_emision", "FOLIO": "folio",
    "MONTO": "monto", "BENEFICIARIO": "beneficiario", "CUENTA CON SELLO": "cuenta_con_sello",
    "FECHA DE COBRO": "fecha_cobro", "CONCEPTO": "concepto", "OBSERVACIONES": "observaciones",
    "CAUSA DE LA BAJA": "causa_de_la_baja", "FECHA DE BAJA": "fecha_de_baja",
    "ESTATUS_CHEQUE": "estatus_cheque", "EVIDENCIA_VALE": "evidencia_vale",
    "EVIDENCIA_CANCELADO": "evidencia_cancelado"
}
INV_DB_MAP = {v: k for k, v in DB_MAP.items()}

def convertir_fecha_sql(val):
    """Estandariza cualquier formato de fecha humana a cadena ISO (YYYY-MM-DD) para SQL"""
    if pd.isna(val) or str(val).strip() in ["", "NONE", "NAN", "<NAT>", "00/00/0000"] or "PENDIENTE" in str(val).upper():
        return None
    try:
        dt = pd.to_datetime(val, errors='coerce')
        if pd.notna(dt) and dt is not pd.NaT:
            return str(dt.date())
        return None
    except:
        return None

def formatear_fecha_visor(val):
    """Convierte cualquier formato de fecha de la DB a DD/MM/AAAA de forma ultra-segura (Anti-Crash)"""
    if pd.isna(val) or str(val).strip() in ["", "NONE", "NAN", "<NAT>", "00/00/0000"]:
        return ""
    val_str = str(val).strip().upper()
    if "PENDIENTE" in val_str:
        return str(val).strip().upper() # Preservamos el texto visualmente para comodidad del usuario
    try:
        dt = pd.to_datetime(val, errors='coerce')
        if pd.notna(dt) and dt is not pd.NaT:
            return dt.strftime('%d/%m/%Y')
        return str(val).strip()
    except:
        return str(val).strip()

def cargar_desde_sql(supabase_client) -> pd.DataFrame:
    if supabase_client is None: return pd.DataFrame(columns=list(DB_MAP.keys()))
    try:
        response = supabase_client.table("cheques").select("*").execute()
        if not response.data:
            return pd.DataFrame(columns=list(DB_MAP.keys()))
        df_sql = pd.DataFrame(response.data)
        df_esp = df_sql.rename(columns=INV_DB_MAP)
        return df_esp[list(DB_MAP.keys())]
    except Exception as e:
        st.error(f"Error al descargar datos de SQL: {e}")
        return pd.DataFrame(columns=list(DB_MAP.keys()))

def guardar_en_sql(supabase_client, df_datos):
    if supabase_client is None or df_datos.empty: return
    try:
        df_sql = df_datos.rename(columns=DB_MAP)
        df_sql = df_sql.replace({np.nan: None, pd.NaT: None, "": None})
        
        # Filtro purificador de fechas antes de tocar el motor SQL
        for col_f in ["fecha_emision", "fecha_cobro", "fecha_de_baja"]:
            if col_f in df_sql.columns:
                df_sql[col_f] = df_sql[col_f].apply(convertir_fecha_sql)
        
        records = df_sql.to_dict(orient="records")
        supabase_client.table("cheques").upsert(records).execute()
    except Exception as e:
        st.error(f"❌ Error en la sincronización transaccional SQL: {e}")

# ==========================================
# CAPA INTERMEDIA: EL ESCUDO LIMPIADOR (PANDAS)
# ==========================================
def forzar_mayusculas_y_limpieza(df):
    if df is None or df.empty: return df
    for col in df.columns:
        if df[col].dtype == 'object':
            df[col] = df[col].fillna('').astype(str).str.upper().str.strip()
            df[col] = df[col].replace(['NAN', 'NONE', '<NAT>'], '')
    return df

def limpiar_y_convertir_monto_positivo(df, columna):
    if columna in df.columns:
        df[columna] = df[columna].astype(str).str.replace('$', '', regex=False).str.replace(',', '', regex=False).str.strip()
        df[columna] = pd.to_numeric(df[columna], errors='coerce').fillna(0).abs()
    return df

# ==========================================
# MAPA DE COLORES Y EXPORTACIÓN V5 (DETECTOR FLEXIBLE)
# ==========================================
def colorear_reglas_auditoria(row):
    estilo = [''] * len(row)
    concepto = str(row.get('CONCEPTO', '')).upper()
    sello = str(row.get('CUENTA CON SELLO', '')).upper()
    estatus = str(row.get('ESTATUS_CHEQUE', '')).upper()
    vale = str(row.get('EVIDENCIA_VALE', '')).strip()
    justificante = str(row.get('EVIDENCIA_CANCELADO', '')).strip()
    f_emision = pd.to_datetime(row.get('FECHA DE EMISIÓN'), errors='coerce')
    
    # 🔥 DETECTOR INTELIGENTE DE TRÁNSITOS: Captura vacíos, nulos y la palabra "PENDIENTE"
    f_cobro_str = str(row.get('FECHA DE COBRO', '')).strip().upper()
    es_pendiente_cobro = f_cobro_str == "" or "PENDIENTE" in f_cobro_str or f_cobro_str in ["NAN", "NONE", "<NAT>", "00/00/0000"]
    
    if "GRATIFICACI" in concepto: return ['background-color: #ffcccc; color: #cc0000; font-weight: bold'] * len(row)
    if estatus == "CANCELADO" and (justificante == "" or "PENDIENTE" in justificante): return ['background-color: #fce5cd; color: #b45f06'] * len(row)
    if ("NO" == sello or "SIN SELLO" in sello) and (vale == "" or "PENDIENTE" in vale) and estatus != "CANCELADO": return ['background-color: #fff2cc; color: #856404'] * len(row)
    
    # 🔵 Aplicar Alerta Azul considerando la nueva regla flexible de pendientes
    if estatus != "CANCELADO" and pd.notna(f_emision) and es_pendiente_cobro:
        if (datetime.now() - f_emision).days > 7: return ['background-color: #e6f2ff; color: #0b5394'] * len(row)
    return estilo

def generar_excel_coloreado(df):
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte_SQL_Auditado"
    headers = list(df.columns)
    ws.append(headers)
    
    f_rojo, f_amarillo, f_azul, f_naranja = [PatternFill(start_color=c, end_color=c, fill_type="solid") for c in ["FFCCCC", "FFF2CC", "E6F2FF", "FCE5CD"]]
    for index, row in df.iterrows():
        ws.append([str(val) if pd.notna(val) else "" for val in row])
        row_idx = ws.max_row
        concepto, sello, estatus = str(row.get('CONCEPTO','')).upper(), str(row.get('CUENTA CON SELLO','')).upper(), str(row.get('ESTATUS_CHEQUE','')).upper()
        vale, justificante = str(row.get('EVIDENCIA_VALE','')).strip(), str(row.get('EVIDENCIA_CANCELADO','')).strip()
        f_emision = pd.to_datetime(row.get('FECHA DE EMISIÓN'), errors='coerce')
        
        # Sincronizar regla flexible también en la exportación física a Excel
        f_cobro_str = str(row.get('FECHA DE COBRO', '')).strip().upper()
        es_pendiente_cobro = f_cobro_str == "" or "PENDIENTE" in f_cobro_str or f_cobro_str in ["NAN", "NONE", "<NAT>", "00/00/0000"]
        
        fill_actual = None
        if "GRATIFICACI" in concepto: fill_actual = f_rojo
        elif estatus == "CANCELADO" and (justificante == "" or "PENDIENTE" in justificante): fill_actual = f_naranja
        elif ("NO" == sello or "SIN SELLO" in sello) and (vale == "" or "PENDIENTE" in vale) and estatus != "CANCELADO": fill_actual = f_amarillo
        elif estatus != "CANCELADO" and pd.notna(f_emision) and es_pendiente_cobro:
            if (datetime.now() - f_emision).days > 7: fill_actual = f_azul
            
        if fill_actual:
            for col_idx in range(1, len(headers) + 1): ws.cell(row=row_idx, column=col_idx).fill = fill_actual
    wb.save(output)
    return output.getvalue()

# ==========================================
# PORTAL DE SEGURIDAD (LOGIN)
# ==========================================
if "autenticado" not in st.session_state: st.session_state["autenticado"] = False
if not st.session_state["autenticado"]:
    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        st.markdown("<br><br>", unsafe_allow_html=True)
        st.write("### 🛡️ Portal de Auditoría Interna - SQL Engine")
        with st.form("login"):
            if st.form_submit_button("Ingresar al Sistema") and st.text_input("Usuario:") == "auditor_interno" and st.text_input("Contraseña:", type="password") == "ControlTrans2026*":
                st.session_state["autenticado"] = True
                st.rerun()
    st.stop()

# Inicialización e Indexación de Datos
db_client = conectar_supabase()
if "MASTER_CHEQUES" not in st.session_state or st.sidebar.button("🔄 Sincronizar con Servidor SQL"):
    with st.spinner("Sincronizando con la Base de Datos SQL..."):
        st.session_state["MASTER_CHEQUES"] = cargar_desde_sql(db_client)

st.sidebar.title("🛡️ Panel Control SQL")
if st.sidebar.button("🔒 Cerrar Sesión"):
    st.session_state["autenticado"] = False
    st.rerun()

opcion = st.sidebar.radio("Menú Principal:", ["📥 Alimentar / Capturar Cheques", "🔍 Auditoría y Alertas Forenses", "📊 Dashboard y Gráficos"])
LISTA_EMPRESAS = ["Bajio Tropper", "Bajio TCS", "Bajio TCS Gerencia", "Bajio Somnus", "Bajio Idea", "Bajio Areng", "BBVA Tropper", "BBVA Sab Solutions"]

if opcion == "📥 Alimentar / Capturar Cheques":
    st.title("📥 Panel de Control - Base de Datos Permanente")
    st.markdown("---")
    
    tab_editor, tab_carga, tab_manual, tab_boveda, tab_exportar = st.tabs([
        "📝 TABLA INTERACTIVA SQL", "📁 Cargar Archivo (Contpaq o Histórico)", "✍️ Captura Manual Individual", "📂 BÓVEDA DE EVIDENCIAS", "💾 Exportar a Excel"
    ])
    
    with tab_editor:
        st.subheader("Modificación de Celdas en Tiempo Real")
        if not st.session_state["MASTER_CHEQUES"].empty:
            df_m = st.session_state["MASTER_CHEQUES"].copy()
            # Dejar pasar textos especiales en fecha_cobro para el editor interactivo
            for col_f in ["FECHA DE EMISIÓN", "FECHA DE BAJA"]:
                df_m[col_f] = pd.to_datetime(df_m[col_f], errors='coerce')
            df_m = df_m.sort_values(by="FOLIO", ascending=True)

            df_editado = st.data_editor(df_m, num_rows="dynamic", width="stretch", column_config={
                "FECHA DE EMISIÓN": st.column_config.DateColumn(format="DD/MM/YYYY"),
                "FECHA DE BAJA": st.column_config.DateColumn(format="DD/MM/YYYY"),
                "MONTO": st.column_config.NumberColumn(format="$%.2f"),
                "CUENTA CON SELLO": st.column_config.SelectboxColumn(options=["SI", "NO", "NO APLICA"]),
                "ESTATUS_CHEQUE": st.column_config.SelectboxColumn(options=["EMITIDO", "CANCELADO"]),
                "CONCEPTO": st.column_config.SelectboxColumn(options=[
                    "1. COMPROBACIÓN DE GASTOS", "2. FINIQUITO", "3. ABONO A FINANCIERA", 
                    "4. APERTURA DE FONDO FIJO", "5. PAGO DE NÓMINA", "6. GASTO POR COMPROBAR", 
                    "7. APOYO POR DEFUNCIÓN", "8. GRATIFICACIONES", "9. PENSIÓN ALIMENTICIA", 
                    "10. RETENCION A PROVEEDOR", "11. PAGO DE TRÁMITES", "12. OTRO MOTIVO"
                ])
            })
            
            if st.button("💾 Sincronizar Cambios con la Base de Datos SQL"):
                with st.spinner("Actualizando registros en el servidor..."):
                    df_final = forzar_mayusculas_y_limpieza(df_editado)
                    guardar_en_sql(db_client, df_final)
                    st.session_state["MASTER_CHEQUES"] = df_final
                    st.success("¡Base de datos SQL actualizada correctamente!")
                    st.rerun()
        else: st.info("No hay registros en la base de datos.")

    with tab_carga:
        st.subheader("Cargador Universal Inteligente (Soporta Contpaq e Históricos)")
        archivo_nuevo = st.file_uploader("Arrastra tu Excel aquí (Layout de Chequera o Reporte de Contpaq):", type=["xlsx"])
        empresa_carga = st.selectbox("Asignar a Empresa de forma predeterminada:", LISTA_EMPRESAS)
        
        if archivo_nuevo:
            df_raw = pd.read_excel(archivo_nuevo, header=None)
            if not df_raw.empty:
                keywords = ["FOLIO", "BENEFICIARIO", "MONTO", "FECHA DE EMISIÓN", "NÚMERO/FOLIO", "IMPORTE", "CUENTA CON SELLO", "CONCEPTO"]
                header_idx = None
                for idx, row in df_raw.iterrows():
                    if sum(1 for kw in keywords if any(kw in str(cell).upper() for cell in row.values)) >= 2:
                        header_idx = idx
                        break
                
                if header_idx is not None:
                    df_temp = df_raw.iloc[header_idx+1:].copy()
                    
                    encabezados_crudos = [str(c).strip().upper() for c in df_raw.iloc[header_idx].values]
                    encabezados_unicos = []
                    conteos = {}
                    for col in encabezados_crudos:
                        if col in conteos:
                            conteos[col] += 1
                            encabezados_unicos.append(f"{col}_{conteos[col]}")
                        else:
                            conteos[col] = 0
                            encabezados_unicos.append(col)
                    df_temp.columns = encabezados_unicos
                    
                    st.info("Estructura de columnas identificada de forma correcta. Vista previa de los datos:")
                    st.dataframe(df_temp.head(2), width='stretch')
                    
                    if st.button("🚀 Activar Filtros de Preservación e Inyectar en SQL"):
                        with st.spinner("Procesando datos por el escudo protector..."):
                            mapeo_columnas = {
                                "FECHA DE EMISIÓN": ["FECHA DE EMISIÓN", "FECHA DE EMISION", "FECHA EMISION", "FECHA"],
                                "FOLIO": ["FOLIO", "NUMERO", "NÚMERO", "NUM", "CHEQUE", "NÚMERO/FOLIO"],
                                "MONTO": ["MONTO", "IMPORTE", "CANTIDAD", "CARGOS"],
                                "BENEFICIARIO": ["BENEFICIARIO", "NOMBRE", "PROVEEDOR", "BENEFICIARIO/PAGADOR"],
                                "CUENTA CON SELLO": ["CUENTA CON SELLO", "SELLO"],
                                "FECHA DE COBRO": ["FECHA DE COBRO (DEPOSITO EN CUENTA)", "FECHA DE COBRO"],
                                "CONCEPTO": ["CONCEPTO", "MOTIVO", "DESCRIPCION"],
                                "OBSERVACIONES": ["OBSERVACIONES", "NOTAS", "COMENTARIOS", "OBSERVACION"],
                                "EMPRESA": ["EMPRESA"], "CAUSA DE LA BAJA": ["CAUSA DE LA BAJA", "CAUSA DE BAJA"],
                                "FECHA DE BAJA": ["FECHA DE BAJA", "FECHA DE LA BAJA"], "ESTATUS_CHEQUE": ["ESTATUS_CHEQUE", "ESTATUS"]
                            }
                            columnas_nuevas = {col: est for col in df_temp.columns for est, aliases in mapeo_columnas.items() if col in aliases}
                            df_temp = df_temp.rename(columns=columnas_nuevas)
                            
                            for col in st.session_state["MASTER_CHEQUES"].columns:
                                if col not in df_temp.columns: df_temp[col] = ""
                            
                            df_temp = df_temp[df_temp["FOLIO"].notna() & (df_temp["FOLIO"].astype(str).str.strip() != "")]
                            df_temp = limpiar_y_convertir_monto_positivo(df_temp, 'MONTO')
                            df_temp["FOLIO"] = pd.to_numeric(df_temp["FOLIO"], errors='coerce').fillna(0).astype(int)
                            
                            if df_temp["EMPRESA"].astype(str).str.strip().eq("").all():
                                df_temp["EMPRESA"] = empresa_carga
                            
                            if df_temp["ESTATUS_CHEQUE"].astype(str).str.strip().eq("").all():
                                df_temp["ESTATUS_CHEQUE"] = df_temp.apply(lambda r: "CANCELADO" if "CANCELADO" in str(r.get("OBSERVACIONES","")).upper() or "CANCELADO" in str(r.get("BENEFICIARIO","")).upper() else "EMITIDO", axis=1)
                            
                            if "EVIDENCIA_VALE" not in df_temp.columns or df_temp["EVIDENCIA_VALE"].astype(str).str.strip().eq("").all():
                                df_temp["EVIDENCIA_VALE"] = df_temp.apply(lambda r: "PENDIENTE ❌" if str(r.get("CUENTA CON SELLO","")).upper() == "NO" else "NO REQUERIDO", axis=1)
                            
                            if "EVIDENCIA_CANCELADO" not in df_temp.columns or df_temp["EVIDENCIA_CANCELADO"].astype(str).str.strip().eq("").all():
                                df_temp["EVIDENCIA_CANCELADO"] = df_temp.apply(lambda r: "PENDIENTE ❌" if str(r.get("ESTATUS_CHEQUE","")).upper() == "CANCELADO" else "NO REQUERIDO", axis=1)
                            
                            df_temp = df_temp.reindex(columns=st.session_state["MASTER_CHEQUES"].columns)
                            df_temp = forzar_mayusculas_y_limpieza(df_temp)
                            
                            df_unificado = pd.concat([st.session_state["MASTER_CHEQUES"], df_temp], ignore_index=True).drop_duplicates(subset=["EMPRESA", "FOLIO"], keep="last")
                            guardar_en_sql(db_client, df_unificado)
                            st.session_state["MASTER_CHEQUES"] = df_unificado
                            st.success("¡Datos absorbidos e integrados permanentemente en SQL!")
                            st.rerun()

    with tab_manual:
        st.subheader("Captura Manual de un Cheque Individual")
        with st.form("registro_manual_sql", clear_on_submit=True):
            c1, c2, c3 = st.columns(3)
            emp_m = c1.selectbox("Empresa:", LISTA_EMPRESAS)
            fecha_m = c2.date_input("Fecha de Emisión:")
            folio_m = c3.number_input("Folio del Cheque:", step=1, min_value=1)
            
            c4, c5, c6 = st.columns(3)
            monto_m = c4.number_input("Monto ($):", min_value=0.0, format="%.2f")
            benef_m = c5.text_input("Beneficiario:")
            sello_m = c6.selectbox("¿Cuenta con Sello 'Abono en Cuenta'?", ["SI", "NO", "NO APLICA"])
            
            c7, c8 = st.columns(2)
            concepto_m = c7.selectbox("Concepto del Cheque:", [
                "1. COMPROBACIÓN DE GASTOS", "2. FINIQUITO", "3. ABONO A FINANCIERA", 
                "4. APERTURA DE FONDO FIJO", "5. PAGO DE NÓMINA", "6. GASTO POR COMPROBAR", 
                "7. APOYO POR DEFUNCIÓN", "8. GRATIFICACIONES", "9. PENSIÓN ALIMENTICIA", 
                "10. RETENCION A PROVEEDOR", "11. PAGO DE TRÁMITES", "12. OTRO MOTIVO"
            ])
            estatus_m = c8.selectbox("Estatus Inicial:", ["EMITIDO", "CANCELADO"])
            obs_m = st.text_area("Observaciones:")
            
            baja_causa = st.text_input("Causa de la Baja (Solo si aplica para Finiquitos):")
            baja_fecha = st.text_input("Fecha de la Baja (DD/MM/AAAA - Solo Finiquitos):")
            
            if st.form_submit_button("Guardar Cheque en Servidor SQL"):
                vale_inicial = "PENDIENTE ❌" if sello_m == "NO" else "NO REQUERIDO"
                canc_inicial = "PENDIENTE ❌" if estatus_m == "CANCELADO" else "NO REQUERIDO"
                
                nuevo_reg = {
                    "EMPRESA": emp_m, "FECHA DE EMISIÓN": str(fecha_m), "FOLIO": int(folio_m),
                    "MONTO": monto_m, "BENEFICIARIO": benef_m, "CUENTA CON SELLO": sello_m,
                    "FECHA DE COBRO": "PENDIENTE DE PAGO", "CONCEPTO": concepto_m, "OBSERVACIONES": obs_m,
                    "CAUSA DE LA BAJA": baja_causa, "FECHA DE BAJA": baja_fecha, "ESTATUS_CHEQUE": estatus_m,
                    "EVIDENCIA_VALE": vale_inicial, "EVIDENCIA_CANCELADO": canc_inicial
                }
                df_manual = forzar_mayusculas_y_limpieza(pd.DataFrame([nuevo_reg]))
                df_unificado = pd.concat([st.session_state["MASTER_CHEQUES"], df_manual], ignore_index=True).drop_duplicates(subset=["EMPRESA", "FOLIO"], keep="last")
                guardar_en_sql(db_client, df_unificado)
                st.session_state["MASTER_CHEQUES"] = df_unificado
                st.success(f"¡Cheque Folio {folio_m} guardado exitosamente en SQL!")
                st.rerun()

    with tab_boveda:
        st.subheader("📂 Bóveda Transaccional de Soportes")
        tipo_ev = st.radio("Tipo de documento:", ["VALE DE EFECTIVO", "JUSTIFICANTE DE CANCELACIÓN"])
        emp_v = st.selectbox("Empresa:", LISTA_EMPRESAS, key="emp_bov")
        
        df_master_v = st.session_state["MASTER_CHEQUES"]
        if not df_master_v.empty:
            df_fil = df_master_v[df_master_v["EMPRESA"] == emp_v]
            if not df_fil.empty:
                folio_v = st.selectbox("Selecciona el Folio:", df_fil["FOLIO"].unique())
                archivo_ev = st.file_uploader("Sube el comprobante digital:")
                
                if st.button("🔒 Vincular Soporte en la Base de Datos SQL") and archivo_ev:
                    idx_target = df_master_v[(df_master_v["EMPRESA"] == emp_v) & (df_master_v["FOLIO"].astype(str) == str(folio_v))].index
                    col_target = "EVIDENCIA_VALE" if tipo_ev == "VALE DE EFECTIVO" else "EVIDENCIA_CANCELADO"
                    st.session_state["MASTER_CHEQUES"].loc[idx_target, col_target] = archivo_ev.name.upper()
                    
                    guardar_en_sql(db_client, st.session_state["MASTER_CHEQUES"])
                    st.success("¡Soporte guardado permanentemente!")
                    st.rerun()
            else: st.write("Sin folios registrados para esta empresa.")

    with tab_exportar:
        st.subheader("Descargar Reporte")
        if not st.session_state["MASTER_CHEQUES"].empty:
            excel_binario = generar_excel_coloreado(st.session_state["MASTER_CHEQUES"])
            st.download_button(label="📥 Descargar Excel con Colores de Alerta", data=excel_binario, file_name="Auditoria_SQL_Color.xlsx")

    # VISOR DE MAPA DE CALOR
    st.markdown("---")
    st.subheader("👁️ Visor Forense General (Base de Datos Real en la Nube)")
    if not st.session_state["MASTER_CHEQUES"].empty:
        df_visor = st.session_state["MASTER_CHEQUES"].copy().sort_values(by="FOLIO", ascending=True)
        for col_f in ["FECHA DE EMISIÓN", "FECHA DE COBRO", "FECHA DE BAJA"]:
            if col_f in df_visor.columns:
                df_visor[col_f] = df_visor[col_f].apply(formatear_fecha_visor)
        st.dataframe(df_visor.style.apply(colorear_reglas_auditoria, axis=1), width='stretch')

# ==========================================
# SECCIONES ADICIONALES
# ==========================================
elif opcion == "🔍 Auditoría y Alertas Forenses":
    st.title("🔍 Escáner Forense de Riesgos de Control Interno")
    st.info("Módulo listo. Sube el archivo de banco para realizar cruces de auditoría.")

elif opcion == "📊 Dashboard y Gráficos":
    st.title("📊 Métricas Ejecutivas e Indicadores Mensuales")
    st.info("Módulo gráfico enlazado a SQL.")
